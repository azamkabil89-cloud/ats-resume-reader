"""AI Resume ATS Checker - Streamlit + Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# Tried in order; the first one that works is used. Override with the
# GEMINI_MODEL secret / environment variable.
DEFAULT_MODELS = ["gemini-3.6-flash", "gemini-2.5-flash"]
MAX_RESUME_CHARS = 30_000

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume the user provides.

Return ONLY a JSON object with exactly this structure:
{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "category_scores": {
    "formatting": <0-100>,
    "keywords": <0-100>,
    "content_quality": <0-100>,
    "readability": <0-100>
  },
  "strengths": ["..."],
  "weaknesses": ["..."],
  "improvements": [
    {"priority": "High|Medium|Low", "section": "...", "issue": "...", "suggestion": "..."}
  ],
  "missing_keywords": ["..."]
}

Scoring guide: 90+ excellent, 75-89 good, 60-74 average, below 60 needs work.
Judge: standard section headings, contact info, quantified achievements, strong
action verbs, keyword relevance, consistent dates, and clean parsable structure.
If a job description is supplied, judge keyword match against it and list the
important missing keywords. Otherwise judge general ATS-readiness.
Give 5-8 specific, actionable improvements. Be honest, not flattering."""


# ----------------------------- file parsing -----------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------- Gemini helpers -----------------------------
def parse_json_response(text: str) -> dict:
    """Parse model output as JSON, tolerating markdown code fences."""
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip(), flags=re.IGNORECASE)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def normalize_result(raw: dict) -> dict:
    """Fill in defaults and clamp scores so the UI never crashes."""

    def clamp(v):
        try:
            return max(0, min(100, int(round(float(v)))))
        except (TypeError, ValueError):
            return 0

    cats = raw.get("category_scores") or {}
    return {
        "ats_score": clamp(raw.get("ats_score")),
        "summary": str(raw.get("summary", "")),
        "category_scores": {
            k: clamp(cats.get(k))
            for k in ("formatting", "keywords", "content_quality", "readability")
        },
        "strengths": [str(s) for s in raw.get("strengths") or []],
        "weaknesses": [str(s) for s in raw.get("weaknesses") or []],
        "improvements": [i for i in raw.get("improvements") or [] if isinstance(i, dict)],
        "missing_keywords": [str(s) for s in raw.get("missing_keywords") or []],
    }


def analyze_resume(api_key: str, resume_text: str, job_description: str = "") -> dict:
    """Send the resume to Gemini and return the normalized analysis."""
    client = genai.Client(api_key=api_key)
    prompt = f"RESUME:\n{resume_text[:MAX_RESUME_CHARS]}"
    if job_description.strip():
        prompt += f"\n\nTARGET JOB DESCRIPTION:\n{job_description.strip()[:10_000]}"

    custom = os.environ.get("GEMINI_MODEL")
    models = [custom] if custom else DEFAULT_MODELS
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        response_mime_type="application/json",
        temperature=0.2,
    )

    last_error = None
    for model in models:
        try:
            response = client.models.generate_content(
                model=model, contents=prompt, config=config
            )
            return normalize_result(parse_json_response(response.text))
        except Exception as exc:  # try the next model
            last_error = exc
    raise RuntimeError(f"Gemini request failed: {last_error}")


# ----------------------------- UI -----------------------------
def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file locally
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "")


def score_color(score: int) -> str:
    return "🟢" if score >= 75 else "🟡" if score >= 60 else "🔴"


def render_results(result: dict) -> None:
    score = result["ats_score"]
    st.subheader(f"{score_color(score)} ATS Score: {score}/100")
    st.progress(score / 100)
    if result["summary"]:
        st.write(result["summary"])

    cols = st.columns(4)
    labels = {
        "formatting": "Formatting",
        "keywords": "Keywords",
        "content_quality": "Content",
        "readability": "Readability",
    }
    for col, (key, label) in zip(cols, labels.items()):
        col.metric(label, f"{result['category_scores'][key]}/100")

    left, right = st.columns(2)
    with left:
        st.markdown("### ✅ Strengths")
        for s in result["strengths"] or ["None identified."]:
            st.markdown(f"- {s}")
    with right:
        st.markdown("### ⚠️ Weaknesses")
        for s in result["weaknesses"] or ["None identified."]:
            st.markdown(f"- {s}")

    st.markdown("### 🛠️ Suggested Improvements")
    order = {"high": 0, "medium": 1, "low": 2}
    items = sorted(
        result["improvements"],
        key=lambda i: order.get(str(i.get("priority", "")).lower(), 3),
    )
    for item in items:
        prio = str(item.get("priority", "Medium"))
        title = f"[{prio}] {item.get('section', 'General')}"
        with st.expander(title, expanded=prio.lower() == "high"):
            st.markdown(f"**Issue:** {item.get('issue', '')}")
            st.markdown(f"**Fix:** {item.get('suggestion', '')}")

    if result["missing_keywords"]:
        st.markdown("### 🔑 Missing Keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.download_button(
        "Download report (JSON)",
        json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume to get an ATS score and concrete improvements.")

    api_key = get_api_key()
    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password")
            st.markdown("[Get a free key](https://aistudio.google.com/apikey)")
        else:
            st.success("API key loaded")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_desc = st.text_area(
        "Job description (optional, improves keyword matching)", height=150
    )

    if st.button("Analyze Resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        try:
            text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return
        if len(text) < 50:
            st.error(
                "Couldn't extract enough text. If your PDF is a scan or image, "
                "that's also a problem for real ATS systems - use a text-based file."
            )
            return
        with st.spinner("Analyzing your resume..."):
            try:
                st.session_state["result"] = analyze_resume(api_key, text, job_desc)
            except Exception as exc:
                st.error(str(exc))
                return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
